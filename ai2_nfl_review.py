"""Readable and protected NFL Classic portfolio exports.

UI-agnostic; no approvals or contest changes happen during automatic builds.
"""
from collections import Counter
from datetime import datetime, timezone
from io import BytesIO
import json

import pandas as pd

from ai2_nfl_export import SLOTS, slot_classic_lineup, export_classic_candidates

DISPLAY_SLOTS = ("QB", "RB1", "RB2", "WR1", "WR2", "WR3", "TE", "FLEX", "DST")


def review_tables(ranked):
    if not ranked:
        raise ValueError("No portfolio to review")
    overview, detail, exposures = [], [], Counter()
    premium = Counter()
    seen = set()
    for idx, item in enumerate(ranked, 1):
        lineup = item["candidate"]
        roster = slot_classic_lineup(lineup["players"])
        ids = tuple(sorted(str(p["ID"]) for p in roster))
        if ids in seen:
            raise ValueError("Duplicate roster in portfolio")
        seen.add(ids)
        contest = str(item["contest_tier"])
        total = int(sum(int(p["Salary"]) for p in roster))
        if total != int(lineup["salary"]):
            raise ValueError("Salary mismatch in portfolio")
        row = {"Lineup": idx, "Contest": contest, "Salary": total,
               "Remaining": 50000 - total,
               "Mean": round(float(item["simulated_mean"]), 2),
               "P90": round(float(item["simulated_p90"]), 2),
               "P99": round(float(item["simulated_p99"]), 2)}
        for slot, p in zip(DISPLAY_SLOTS, roster):
            pid = str(p["ID"])
            name = str(p["Name"])
            team = str(p.get("TeamAbbrev", ""))
            row[slot] = f"{name} ({team})"
            detail.append({"Lineup": idx, "Contest": contest, "Slot": slot,
                           "Player": name, "Team": team,
                           "Position": p["Position"], "Salary": int(p["Salary"]),
                           "Player mean": round(float(p.get("proj", 0) or 0), 2),
                           "DraftKings ID": pid})
            exposures[(pid, name, team, str(p["Position"]))] += 1
            if contest == "$3":
                premium[(pid, name, team, str(p["Position"]))] += 1
        overview.append(row)
    exposure = [{"Player": name, "Team": team, "Position": pos,
                 "Total / 20": times, "$3 / 5": premium.get((pid, name, team, pos), 0),
                 "DraftKings ID": pid}
                for (pid, name, team, pos), times in exposures.most_common()]
    return (pd.DataFrame(overview), pd.DataFrame(detail), pd.DataFrame(exposure))


def named_csv(ranked):
    overview, _, _ = review_tables(ranked)
    return overview.to_csv(index=False).encode("utf-8-sig")


def named_excel(ranked):
    overview, detail, exposure = review_tables(ranked)
    out = BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        for name, frame in [("20 lineup review", overview),
                            ("Player details", detail),
                            ("Player exposure", exposure)]:
            frame.to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            ws.freeze_panes = "C2"
            ws.auto_filter.ref = ws.dimensions
            from openpyxl.styles import Font, PatternFill
            for cell in ws[1]:
                cell.fill = PatternFill("solid", fgColor="152B48")
                cell.font = Font(bold=True, color="FFFFFF")
            for column in ws.columns:
                letter = column[0].column_letter
                longest = max(len(str(cell.value or "")) for cell in column[:min(len(column), 100)])
                ws.column_dimensions[letter].width = min(43, max(12, longest + 3))
    return out.getvalue()


def swap_contest_tiers(ranked, premium_lineup, regular_lineup):
    """Swap stakes, without changing any roster or the original lineup order."""
    import copy
    if not 1 <= premium_lineup <= len(ranked) or not 1 <= regular_lineup <= len(ranked):
        raise ValueError("Lineup number out of range")
    new = copy.deepcopy(ranked)
    i, j = premium_lineup - 1, regular_lineup - 1
    if new[i]["contest_tier"] != "$3" or new[j]["contest_tier"] != "$0.50":
        raise ValueError("Choose one current $3 lineup and one current $0.50 lineup")
    new[i]["contest_tier"], new[j]["contest_tier"] = "$0.50", "$3"
    return new


def snapshot_approved(ranked, slate_digest, seed):
    """A portable human-approved copy, distinct from the next automatic build."""
    overview, _, _ = review_tables(ranked)
    safe_ranked = []
    for item in ranked:
        safe_ranked.append({
            "contest_tier": item["contest_tier"],
            "simulated_mean": float(item["simulated_mean"]),
            "simulated_p90": float(item["simulated_p90"]),
            "simulated_p99": float(item["simulated_p99"]),
            "candidate": {
                "salary": int(item["candidate"]["salary"]),
                "players": [{
                    "ID": str(p["ID"]), "Name": str(p["Name"]),
                    "TeamAbbrev": str(p.get("TeamAbbrev", "")),
                    "Position": str(p["Position"]), "Salary": int(p["Salary"]),
                    "proj": float(p.get("proj", 0) or 0)}
                    for p in item["candidate"]["players"]]}})
    return json.dumps({
        "format": "ai2-nfl-approved-v1",
        "approved_utc": datetime.now(timezone.utc).isoformat(),
        "slate_sha256": slate_digest,
        "simulation_seed": int(seed),
        "lineups": safe_ranked,
        "premium_lineup_numbers": overview.loc[overview["Contest"] == "$3", "Lineup"].tolist()
    }, indent=2).encode("utf-8")


def load_approved(payload):
    value = json.loads(payload)
    if value.get("format") != "ai2-nfl-approved-v1" or not value.get("lineups"):
        raise ValueError("Unsupported approved portfolio snapshot")
    review_tables(value["lineups"])
    return value


def roster_id_csv(ranked):
    return export_classic_candidates([item["candidate"] for item in ranked])
